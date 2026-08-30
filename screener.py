"""
Pullback-in-Uptrend Stock Screener
-----------------------------------
Finds stocks that are in a confirmed uptrend AND pulling back to a
healthy entry zone -- while trying to avoid stocks that are actually
rolling over from an uptrend into a downtrend.

Run manually:
    pip install -r requirements.txt
    python screener.py

Or let the included GitHub Actions workflow run it on a schedule and
commit the updated Output/screener_results.xlsx back to the repo.
"""

import pandas as pd
import yfinance as yf
from datetime import datetime
import os
import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.application import MIMEApplication

# ---------------------------------------------------------------------------
# CONFIG -- edit this section to change strictness. Tickers now live in
# config/tickers.txt -- edit that file (one ticker per line) to add/remove
# stocks without touching this script.
# ---------------------------------------------------------------------------

TICKERS_FILE = os.path.join("config", "tickers.txt")

CONFIG = {
    # --- Core trend + pullback (always applied) ---
    "rsi_pullback_max": 45,       # RSI below this = "pulled back"
    "require_above_sma200": True, # price must still be above the 200-day

    # --- Filter 1: SMA50 still rising (trend strength) ---
    "check_sma50_rising": True,
    "sma50_lookback_days": 20,    # compare SMA50 now vs. this many days ago

    # --- Filter 2: Volume confirmation (pullback on lower volume = healthy) ---
    "check_volume_contraction": True,
    "volume_recent_days": 5,      # recent volume window
    "volume_baseline_days": 20,   # baseline volume window
    "volume_ratio_max": 1.1,      # recent avg vol must be <= this x baseline
                                   # (moderate: allow slightly elevated, not a spike)

    # --- Filter 3: Cap pullback depth (avoid deep breakdowns) ---
    "check_pullback_depth": True,
    "lookback_high_days": 63,     # ~3 months of trading days
    "max_pullback_from_high_pct": 15,  # disqualify if down >15% from that high

    # --- Filter 4: MACD reversal check ---
    # A short-term MACD-below-signal dip is NORMAL during any healthy
    # pullback, so that alone must not disqualify a stock. What actually
    # signals "this uptrend is over" is the MACD line dropping below the
    # zero line -- meaning the 12-day trend has genuinely gone negative
    # relative to the 26-day trend, not just short-term noise.
    "check_macd": True,
    "macd_fast": 12,
    "macd_slow": 26,
    "macd_signal": 9,
}

OUTPUT_DIR = "Output"
OUTPUT_FILE = os.path.join(OUTPUT_DIR, "screener_results.xlsx")

# ---------------------------------------------------------------------------
# Email / SMS delivery -- all optional, controlled entirely by environment
# variables (set as GitHub Actions secrets, never hardcoded here).
#
#   SMTP_SERVER    e.g. smtp.gmail.com
#   SMTP_PORT      e.g. 587
#   SMTP_USERNAME  the sending account's email address
#   SMTP_PASSWORD  an app password (NOT your regular account password)
#   EMAIL_TO       where to send the full formatted table (comma-separated
#                  for multiple recipients)
#   SMS_TO         optional: a carrier email-to-SMS gateway address, e.g.
#                  5551234567@vtext.com (Verizon), @txt.att.net (AT&T),
#                  @tmomail.net (T-Mobile). Gets a short PASS-only summary,
#                  not the full table -- SMS can't render a table.
#
# If SMTP_SERVER/USERNAME/PASSWORD aren't all set, email sending is
# silently skipped (so local/manual runs without email configured still
# work fine).
# ---------------------------------------------------------------------------

# ---------------------------------------------------------------------------
# Ticker list
# ---------------------------------------------------------------------------

def load_tickers(path=TICKERS_FILE):
    if not os.path.exists(path):
        raise FileNotFoundError(
            f"Ticker list not found at {path}. Create it with one ticker "
            f"per line (lines starting with # are ignored)."
        )
    tickers = []
    with open(path, "r") as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#"):
                continue
            tickers.append(line.upper())
    if not tickers:
        raise ValueError(f"{path} exists but contains no tickers.")
    return tickers


# ---------------------------------------------------------------------------
# Indicator calculations
# ---------------------------------------------------------------------------

def compute_rsi(close, period=14):
    delta = close.diff()
    gain = delta.where(delta > 0, 0).rolling(period).mean()
    loss = (-delta.where(delta < 0, 0)).rolling(period).mean()
    rs = gain / loss
    return 100 - (100 / (1 + rs))


def compute_macd(close, fast=12, slow=26, signal=9):
    ema_fast = close.ewm(span=fast, adjust=False).mean()
    ema_slow = close.ewm(span=slow, adjust=False).mean()
    macd_line = ema_fast - ema_slow
    signal_line = macd_line.ewm(span=signal, adjust=False).mean()
    return macd_line, signal_line


def evaluate_ticker(symbol, cfg):
    df = yf.download(symbol, period="1y", progress=False, auto_adjust=True)
    if df.empty or len(df) < 210:
        return None

    # yfinance sometimes returns MultiIndex columns for a single ticker
    if isinstance(df.columns, pd.MultiIndex):
        df.columns = df.columns.get_level_values(0)

    df['SMA50'] = df['Close'].rolling(50).mean()
    df['SMA200'] = df['Close'].rolling(200).mean()
    df['RSI14'] = compute_rsi(df['Close'], 14)
    macd_line, signal_line = compute_macd(
        df['Close'], cfg["macd_fast"], cfg["macd_slow"], cfg["macd_signal"]
    )
    df['MACD'] = macd_line
    df['MACDSignal'] = signal_line

    latest = df.iloc[-1]
    close = float(latest['Close'])
    sma50 = float(latest['SMA50'])
    sma200 = float(latest['SMA200'])
    rsi14 = float(latest['RSI14'])

    reasons_failed = []

    # --- Core: uptrend structure ---
    if not (sma50 > sma200):
        reasons_failed.append("no golden cross (SMA50 <= SMA200)")
    if cfg["require_above_sma200"] and not (close > sma200):
        reasons_failed.append("price below SMA200")
    if not (rsi14 < cfg["rsi_pullback_max"]):
        reasons_failed.append(f"RSI not pulled back (RSI={rsi14:.1f})")

    # --- Filter 1: SMA50 still rising ---
    sma50_rising = None
    if cfg["check_sma50_rising"]:
        lb = cfg["sma50_lookback_days"]
        if len(df) > lb and not pd.isna(df['SMA50'].iloc[-lb - 1]):
            sma50_prior = float(df['SMA50'].iloc[-lb - 1])
            sma50_rising = sma50 > sma50_prior
            if not sma50_rising:
                reasons_failed.append("SMA50 flattening/falling (trend weakening)")

    # --- Filter 2: Volume contraction on the pullback ---
    volume_ok = None
    if cfg["check_volume_contraction"]:
        recent_vol = df['Volume'].iloc[-cfg["volume_recent_days"]:].mean()
        baseline_vol = df['Volume'].iloc[-cfg["volume_baseline_days"]:].mean()
        if baseline_vol > 0:
            ratio = recent_vol / baseline_vol
            volume_ok = ratio <= cfg["volume_ratio_max"]
            if not volume_ok:
                reasons_failed.append(f"volume elevated on pullback (ratio={ratio:.2f})")

    # --- Filter 3: Pullback depth cap ---
    depth_ok = None
    pullback_pct = None
    if cfg["check_pullback_depth"]:
        recent_high = df['Close'].iloc[-cfg["lookback_high_days"]:].max()
        pullback_pct = (recent_high - close) / recent_high * 100
        depth_ok = pullback_pct <= cfg["max_pullback_from_high_pct"]
        if not depth_ok:
            reasons_failed.append(f"pullback too deep ({pullback_pct:.1f}% off high)")

    # --- Filter 4: MACD reversal check ---
    # Hard disqualifier: MACD line below zero (12-day trend has genuinely
    # turned negative vs. the 26-day trend -- this is what a real
    # rollover looks like, not a normal dip).
    # Informational only: whether MACD is currently below its signal
    # line -- expected during any pullback, so it's reported but does
    # NOT fail the screen on its own.
    macd_ok = None
    macd_below_signal = None
    if cfg["check_macd"]:
        macd_val = float(latest['MACD'])
        signal_val = float(latest['MACDSignal'])
        macd_below_signal = macd_val < signal_val
        macd_ok = macd_val > 0
        if not macd_ok:
            reasons_failed.append(f"MACD below zero line (momentum turned negative, MACD={macd_val:.2f})")

    passed = len(reasons_failed) == 0

    return {
        "Ticker": symbol,
        "Result": "PASS" if passed else "FAIL",
        "Price": round(close, 2),
        "SMA 50": round(sma50, 2),
        "SMA 200": round(sma200, 2),
        "RSI (14)": round(rsi14, 2),
        "SMA50 Rising": sma50_rising,
        "Volume OK": volume_ok,
        "Pullback % from High": round(pullback_pct, 1) if pullback_pct is not None else None,
        "MACD OK": macd_ok,
        "MACD Below Signal (normal in dip)": macd_below_signal,
        "Reasons Failed": "; ".join(reasons_failed) if reasons_failed else "",
    }


def build_text_table(df_out):
    """Fixed-width plain-text table, PASS rows first. Renders cleanly in
    monospace email clients and in the GitHub Actions log."""
    if df_out.empty:
        return "No results."

    cols = ["Ticker", "Result", "Price", "RSI (14)", "Pullback % from High"]
    cols = [c for c in cols if c in df_out.columns]
    sort_key = df_out["Result"].map({"PASS": 0, "FAIL": 1, "ERROR": 2}).fillna(3)
    df_sorted = df_out.assign(_sort=sort_key).sort_values("_sort").drop(columns="_sort")

    widths = {c: max(len(c), df_sorted[c].astype(str).map(len).max()) + 2 for c in cols}
    header = "".join(c.ljust(widths[c]) for c in cols)
    sep = "-" * len(header)
    lines = [header, sep]
    for _, row in df_sorted.iterrows():
        lines.append("".join(str(row[c]).ljust(widths[c]) for c in cols))
    return "\n".join(lines)


def build_sms_summary(df_out, max_len=300):
    """Short PASS-only summary suitable for an SMS gateway. No table --
    just a compact comma list, truncated to a safe length."""
    passed = df_out.loc[df_out["Result"] == "PASS", "Ticker"].tolist() if "Result" in df_out else []
    if not passed:
        return f"Screener {datetime.now().strftime('%Y-%m-%d')}: no matches today."
    body = f"Screener {datetime.now().strftime('%Y-%m-%d')} PASS: " + ", ".join(passed)
    if len(body) > max_len:
        body = body[:max_len - 3] + "..."
    return body


def send_email(subject, body_text, to_addrs, attachment_path=None):
    server = os.environ.get("SMTP_SERVER")
    port = os.environ.get("SMTP_PORT", "587")
    username = os.environ.get("SMTP_USERNAME")
    password = os.environ.get("SMTP_PASSWORD")

    if not (server and username and password and to_addrs):
        print("Email not configured (missing SMTP_SERVER/USERNAME/PASSWORD or recipient) -- skipping.")
        return

    msg = MIMEMultipart()
    msg["From"] = username
    msg["To"] = to_addrs
    msg["Subject"] = subject
    msg.attach(MIMEText(body_text, "plain"))

    if attachment_path and os.path.exists(attachment_path):
        with open(attachment_path, "rb") as f:
            part = MIMEApplication(f.read(), Name=os.path.basename(attachment_path))
        part["Content-Disposition"] = f'attachment; filename="{os.path.basename(attachment_path)}"'
        msg.attach(part)

    try:
        with smtplib.SMTP(server, int(port)) as smtp:
            smtp.starttls()
            smtp.login(username, password)
            smtp.sendmail(username, [a.strip() for a in to_addrs.split(",")], msg.as_string())
        print(f"Email sent to {to_addrs}.")
    except Exception as e:
        print(f"Email send failed: {e}")


def run_screener():
    tickers = load_tickers()
    results = []
    for symbol in tickers:
        try:
            row = evaluate_ticker(symbol, CONFIG)
            if row:
                results.append(row)
        except Exception as e:
            results.append({
                "Ticker": symbol,
                "Result": "ERROR",
                "Reasons Failed": str(e),
            })

    df_out = pd.DataFrame(results)

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        df_out.to_excel(writer, sheet_name="Screener", index=False)
        ws = writer.sheets["Screener"]

        # Basic column widths
        widths = {
            "A": 10, "B": 10, "C": 10, "D": 10, "E": 10, "F": 12,
            "G": 14, "H": 12, "I": 20, "J": 10, "K": 45,
        }
        for col, w in widths.items():
            ws.column_dimensions[col].width = w

        # Highlight PASS rows green, FAIL rows light red
        from openpyxl.styles import PatternFill
        green = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        red = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        result_col_idx = list(df_out.columns).index("Result") + 1
        for row_idx in range(2, len(df_out) + 2):
            cell_value = ws.cell(row=row_idx, column=result_col_idx).value
            fill = green if cell_value == "PASS" else (red if cell_value == "FAIL" else None)
            if fill:
                for col_idx in range(1, len(df_out.columns) + 1):
                    ws.cell(row=row_idx, column=col_idx).fill = fill

    n_pass = (df_out["Result"] == "PASS").sum() if "Result" in df_out else 0
    print(f"Scan complete at {datetime.now().isoformat()}. "
          f"{n_pass} of {len(tickers)} tickers passed all filters. "
          f"Results written to {OUTPUT_FILE}")

    table_text = build_text_table(df_out)
    print("\n" + table_text)

    # Full table -> email, with the xlsx attached
    email_to = os.environ.get("EMAIL_TO")
    if email_to:
        subject = f"Stock Screener Results - {datetime.now().strftime('%Y-%m-%d')} ({n_pass} match{'es' if n_pass != 1 else ''})"
        send_email(subject, table_text, email_to, attachment_path=OUTPUT_FILE)

    # Short PASS-only summary -> SMS gateway (no table, no attachment)
    sms_to = os.environ.get("SMS_TO")
    if sms_to:
        sms_body = build_sms_summary(df_out)
        send_email("", sms_body, sms_to)  # most carrier gateways ignore subject


if __name__ == "__main__":
    run_screener()
